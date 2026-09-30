from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Styling
NAVY = RGBColor(15, 23, 42)
CRIMSON = RGBColor(190, 18, 60)
DARK_GRAY = RGBColor(51, 65, 85)
LIGHT_BG = RGBColor(248, 250, 252)

slides_data = [
    {
        "category": "FINAL YEAR PROJECT DEFENSE",
        "title": "A Multi-Modal Deep Learning Framework for Vehicle Theft Detection Using Object Detection",
        "bullets": [
            "Student Name: Alabi Joshua Ibukunoluwa (Matric: 220903035)",
            "Supervisor: Dr. O. A. Jongbo",
            "Department of Computer Science | Faculty of Science",
            "Institution: Ekiti State University, Ado-Ekiti, Nigeria (September 2026)"
        ]
    },
    {
        "category": "OVERVIEW",
        "title": "Presentation Outline",
        "bullets": [
            "Background of the Study & Problem Statement",
            "Research Aim, Specific Objectives & Key Questions",
            "Literature Review, Related Works & Identified Gaps",
            "Proposed Multi-Modal System Architecture & Algorithms",
            "System Implementation (FastAPI, React, YOLOv8, OCR, SFace)",
            "Experimental Results, Latency Benchmarks & Discussion",
            "Academic Contributions, Honest Limitations & Future Work",
            "Conclusion & Oral Defense Question Session"
        ]
    },
    {
        "category": "PROBLEM CONTEXT",
        "title": "Background of the Study",
        "bullets": [
            "Severe Theft Incidence: Over 1.58 million vehicle theft cases reported in Nigeria (2023–2024).",
            "Critical Recovery Deficit: NPF recovered only 1,519 vehicles in 2024—a recovery rate of just 0.09%.",
            "Failure of Conventional Systems: Steering locks, GPS trackers, and alarms operate reactively post-theft.",
            "Proactive Vision Solution: Multi-modal AI combines YOLOv8 object detection, OCR, and face verification."
        ]
    },
    {
        "category": "PROBLEM DEFINITION",
        "title": "Statement of the Problem",
        "bullets": [
            "Inadequacy of Reactive Systems: Tracking gear does not stop theft or identify perpetrators in real time.",
            "Human Surveillance Fatigue: Operators miss up to 90% of CCTV security events after 20 minutes.",
            "Lack of Real-Time Owner Verification: Cameras fail to cross-reference drivers against owner biometrics.",
            "High False Positive Rates: Basic motion sensors lack contextual intelligence, causing nuisance alarms."
        ]
    },
    {
        "category": "RESEARCH GOALS",
        "title": "Aim and Objectives",
        "bullets": [
            "Primary Aim: To design, implement, and evaluate an automated multi-modal vehicle theft detection system.",
            "Objective 1: Design an integrated architecture combining vehicle, plate, face, and database modules.",
            "Objective 2: Compile and preprocess robust image datasets for license plates and owner facial galleries.",
            "Objective 3: Train and optimize a dedicated YOLOv8 plate detector and integrate pretrained vision models.",
            "Objective 4: Evaluate pipeline performance using Precision, Recall, F1-score, mAP, and CPU latency."
        ]
    },
    {
        "category": "LITERATURE REVIEW",
        "title": "Existing Approaches & Identified Gaps",
        "bullets": [
            "Hashmi et al. (2019): ANPR plate matching; failed in low light and lacked face recognition.",
            "Okonkwo et al. (2021): IoT + GPS alerts; purely reactive tracking without visual evidence.",
            "Bello et al. (2022): YOLOv5 parking security; limited to vehicle counting without owner verification.",
            "Adeyemi & Afolabi (2023): GPS + CCTV alerts; lacked deep learning face recognition for intruders.",
            "Research Gap Addressed: Unifying YOLOv8 plate detection, EasyOCR, YuNet/SFace, and rule-based risk scoring."
        ]
    },
    {
        "category": "SYSTEM OVERVIEW",
        "title": "Proposed System Framework",
        "bullets": [
            "Multi-Modal Pipeline: Ingests still images or streams, processing spatial, textual, and biometric data.",
            "Visual Perception Stage: YOLOv8n identifies vehicles, license plates, and nearby human occupants.",
            "Text & Biometric Extraction: EasyOCR extracts plate numbers; YuNet + SFace computes face embeddings.",
            "Centralized Threat Engine: Evaluates combined evidence against configurable security rules (0–100 score).",
            "Automated Owner Alerting: Dispatches instant Twilio SMS and SMTP email alerts with photographic proof."
        ]
    },
    {
        "category": "SYSTEM DESIGN",
        "title": "System Architecture",
        "bullets": [
            "Data Sources: Image uploads (.jpg/.png), video files, and live camera RTSP streams.",
            "FastAPI Backend: Threaded pipeline handling object association, OCR normalization, and vector matching.",
            "Persistence Layer: SQLite database via SQLAlchemy storing vehicle records, face galleries, and audit logs.",
            "User Dashboard: React + Vite + Tailwind interface for live monitoring, registration, and alert reviews."
        ]
    },
    {
        "category": "METHODOLOGY",
        "title": "System Workflow & Algorithm",
        "bullets": [
            "Step 1: Receive input image/frame and detect vehicles, license plates, and persons using YOLOv8n.",
            "Step 2: Associate detected persons and license plate bounding boxes with the nearest vehicle region.",
            "Step 3: Crop plate image -> Apply CLAHE/grayscale preprocessing -> Execute EasyOCR -> Normalize text.",
            "Step 4: Extract 128D SFace embedding from detected faces and calculate Cosine Similarity against owner gallery.",
            "Step 5: Run TheftEngine rules -> Calculate Threat Score (0–100) -> If score >= 60, trigger SMS/Email alerts."
        ]
    },
    {
        "category": "EXPERIMENTAL SETUP",
        "title": "Dataset & Model Configuration",
        "bullets": [
            "Roboflow LPD License Plate Dataset: 1,000 total images (700 Train, 200 Validation, 100 Isolated Test).",
            "Pretrained Vision Models: COCO-pretrained YOLOv8n for vehicles/persons; YuNet & SFace for biometrics.",
            "Training Parameters: Fine-tuned plate detector for 20 epochs, batch size = 4, image size = 416x416 on CPU.",
            "Data Augmentation: HSV color shift, scaling, translation, shear, perspective transform, and mosaic."
        ]
    },
    {
        "category": "EXPERIMENTAL RESULTS",
        "title": "License Plate Detector Training Behavior",
        "bullets": [
            "Total Training Duration: 6,541.6 seconds (~1 hour, 49 minutes on CPU).",
            "Learning Rate Schedule: Warm-up over first 3 epochs (0.00066 -> 0.0018), decaying to 0.00012 by epoch 20.",
            "Optimization Objective: Combined bounding-box regression loss, classification loss, and distribution focal loss.",
            "Checkpoint Selection: Validation performance evaluated every epoch; best weights restored for testing."
        ]
    },
    {
        "category": "SOFTWARE ARCHITECTURE",
        "title": "System Implementation & Stack",
        "bullets": [
            "Backend Framework: FastAPI (Python 3.12) providing REST endpoints and MJPEG video streaming.",
            "Computer Vision Libraries: OpenCV, PyTorch, Ultralytics YOLOv8, and EasyOCR.",
            "Database Layer: SQLite with SQLAlchemy ORM handling relational mapping for vehicles and logs.",
            "Frontend Application: React 18 dashboard built with Vite and Tailwind CSS.",
            "Notification Services: Twilio API for SMS delivery and SMTP integration for rich HTML email alerts."
        ]
    },
    {
        "category": "QUANTITATIVE EVALUATION",
        "title": "Held-Out Test Results (Plate Detector)",
        "bullets": [
            "Precision: 99.68% (0.9968) - Almost zero false-positive license plate detections.",
            "Recall: 100.00% (1.0000) - Perfect detection rate; no annotated test plates were missed.",
            "F1-Score: 99.84% (0.9984) - Exceptional balance between precision and recall.",
            "mAP@0.5: 99.50% (0.9950) - Highly reliable plate localization at 50% IoU threshold.",
            "mAP@0.5:0.95: 81.87% (0.8187) - Reflects stricter bounding-box alignment requirements."
        ]
    },
    {
        "category": "LATENCY ANALYSIS",
        "title": "Runtime Performance Benchmark",
        "bullets": [
            "Hardware Environment: Tested over 50 execution runs on standard CPU hardware.",
            "Mean Inference Latency: 199.43 ms per frame.",
            "Median Latency: 188.35 ms | 95th Percentile Latency: 281.55 ms.",
            "Throughput Benchmark: 5.01 Frames Per Second (FPS).",
            "Deployment Suitability: Fully adequate for still-image verification and sampled frame monitoring."
        ]
    },
    {
        "category": "DECISION ENGINE",
        "title": "Explainable Threat Scoring Rules",
        "bullets": [
            "Recognized Registered Owner: -50 Points (Suppression rule).",
            "Unregistered License Plate: +35 Points | Unknown Person Near Vehicle: +40 Points.",
            "Recognized Person is NOT Owner: +45 Points | Unreadable License Plate: +10 Points.",
            "Crowd Presence (>= 3 Persons): +10 Points | Nighttime Suspicious Activity: +5 Points.",
            "Reported Stolen Flag: Overrides Threat Score directly to 100 (Critical Alert Threshold >= 60)."
        ]
    },
    {
        "category": "ANALYTICAL DISCUSSION",
        "title": "Discussion of Findings",
        "bullets": [
            "High Recall Importance: Achieving 100% recall prevents pipeline failure in downstream OCR/Face steps.",
            "mAP Discrepancy: Drop to 81.87% at mAP@0.5:0.95 due to reduced input resolution (416x416) on small plates.",
            "Multi-Modal Resilience: Threat Engine prevents false alarms by weighing multiple evidence streams.",
            "CPU vs. GPU Realities: 5.01 FPS supports image uploads, but high-frame video requires GPU acceleration."
        ]
    },
    {
        "category": "RESEARCH IMPACT",
        "title": "Contributions to Knowledge",
        "bullets": [
            "Technical Contribution: Fine-tuned compact YOLOv8 plate detector achieving >99% mAP@0.5 on CPU.",
            "System Contribution: Integrated end-to-end framework combining detection, OCR, face biometrics, and web UI.",
            "Practical Contribution: Developed an explainable, rule-based threat engine replacing black-box AI decisions.",
            "Security Impact: Empowers vehicle owners with real-time photographic evidence of unauthorized attempts."
        ]
    },
    {
        "category": "HONEST LIMITATIONS",
        "title": "Academic Limitations",
        "bullets": [
            "Dataset Scope: Plate detector evaluated on 100 isolated test images from a single dataset source.",
            "Module Scope: Quantitative benchmark focused on plate detection; OCR and face verification were functionally validated.",
            "Hardware Limits: Benchmarking conducted on CPU hardware at 416x416 resolution.",
            "Environmental Constraints: Performance degrades under heavy plate occlusion or extreme low-light blur."
        ]
    },
    {
        "category": "CLOSING SUMMARY",
        "title": "Conclusion and Recommendations",
        "bullets": [
            "Conclusion: Proved technical feasibility of a multi-modal vehicle theft detection system with 99.50% mAP@0.5.",
            "Recommendation 1: Deploy on CUDA GPU hardware for high-speed, 30+ FPS video surveillance.",
            "Recommendation 2: Establish strict data governance and encryption for stored biometric face embeddings.",
            "Recommendation 3: Expand testing across diverse regional weather, lighting, and license plate styles."
        ]
    },
    {
        "category": "ACKNOWLEDGEMENT",
        "title": "Thank You / Defense Questions",
        "bullets": [
            "A Multi-Modal Deep Learning Framework for Vehicle Theft Detection Using Object Detection",
            "Student: Alabi Joshua Ibukunoluwa | Supervisor: Dr. O. A. Jongbo",
            "Department of Computer Science | Ekiti State University",
            "Questions, Comments, and Panel Discussion are Welcome."
        ]
    }
]

for item in slides_data:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT_BG
    
    # Text box for header
    tb_header = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0))
    tf_h = tb_header.text_frame
    tf_h.word_wrap = True
    
    p_cat = tf_h.paragraphs[0]
    p_cat.text = item["category"]
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = CRIMSON
    
    p_title = tf_h.add_paragraph()
    p_title.text = item["title"]
    p_title.font.size = Pt(20)
    p_title.font.bold = True
    p_title.font.color.rgb = NAVY
    
    # Text box for bullets
    tb_body = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf_b = tb_body.text_frame
    tf_b.word_wrap = True
    
    for i, bullet in enumerate(item["bullets"]):
        p = tf_b.paragraphs[0] if i == 0 else tf_b.add_paragraph()
        p.text = bullet
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(12)

prs.save("Vehicle_Theft_Detection_Defense.pptx")
print("PowerPoint saved successfully as 'Vehicle_Theft_Detection_Defense.pptx'.")